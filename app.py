#!/usr/bin/env python3
"""PricePilot AI / Wissen assessment workspace. Local dashboard + JWT auth + price alerts."""
from __future__ import annotations
import base64, csv, hashlib, hmac, io, json, os, secrets, time, urllib.request, urllib.error
from collections import defaultdict
from datetime import date, datetime
from email import policy
from email.parser import BytesParser
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse, parse_qs
from alerts import InputError, Price, read_prices, read_rules, find_alerts, format_alert

ROOT = Path(__file__).resolve().parent
FRONTEND = ROOT / 'frontend' / 'index.html'
SAMPLE_PRICES = ROOT / 'data' / 'prices.csv'
SAMPLE_RULES = ROOT / 'rules.json'
DATASETS = {
    'amazon': ROOT / 'data' / 'amazon_daily_prices.csv',
    'online-retail': ROOT / 'data' / 'online_retail_daily_prices.csv',
}
ADMIN_USERNAME = os.getenv('ADMIN_USERNAME', 'admin')
ADMIN_PASSWORD = os.getenv('ADMIN_PASSWORD', 'Admin@12345')
JWT_SECRET = os.getenv('JWT_SECRET', 'CHANGE-ME-local-development-secret')
TOKEN_TTL = 8 * 60 * 60
ACTIVE_DATASETS: dict[str, dict] = {}


def b64u(data: bytes) -> str:
    return base64.urlsafe_b64encode(data).rstrip(b'=').decode('ascii')

def make_token(username: str) -> str:
    now = int(time.time())
    head = b64u(json.dumps({'alg':'HS256','typ':'JWT'}, separators=(',',':')).encode())
    body = b64u(json.dumps({'sub':username,'role':'admin','iat':now,'exp':now+TOKEN_TTL}, separators=(',',':')).encode())
    signing = f'{head}.{body}'.encode()
    sig = b64u(hmac.new(JWT_SECRET.encode(), signing, hashlib.sha256).digest())
    return f'{head}.{body}.{sig}'

def verify_token(token: str) -> dict | None:
    try:
        head, body, signature = token.split('.')
        signing=f'{head}.{body}'.encode()
        expected=b64u(hmac.new(JWT_SECRET.encode(), signing, hashlib.sha256).digest())
        if not hmac.compare_digest(signature, expected): return None
        payload=json.loads(base64.urlsafe_b64decode(body+'='*((4-len(body)%4)%4)))
        if payload.get('exp',0) < int(time.time()) or payload.get('role') != 'admin': return None
        return payload
    except Exception:
        return None

def parse_csv_prices(text: str) -> tuple[str, dict]:
    """Accept standard alert CSV or map supported sales datasets into daily prices."""
    try: rows=list(csv.DictReader(io.StringIO(text.lstrip('\ufeff'))))
    except Exception as exc: raise InputError(f'Could not parse CSV: {exc}') from exc
    if not rows: raise InputError('The uploaded CSV has no data rows.')
    cols=set(rows[0].keys())
    if {'date','instrument','close'}.issubset(cols):
        # Validate and preserve standard input as-is.
        return text, {'format':'Price Alert CSV','rows':len(rows),'instruments':len({r.get('instrument','') for r in rows})}
    # Amazon marketplace export mapping.
    if {'OrderDate','ProductID','UnitPrice'}.issubset(cols):
        grouped=defaultdict(list)
        for row in rows:
            day=_parse_day(row.get('OrderDate')); sku=(row.get('ProductID') or '').strip()
            price=_number(row.get('UnitPrice'))
            if day and sku and price is not None and price >= 0: grouped[(day,'AMZ-'+sku)].append(price)
        return _grouped_csv(grouped, 'Amazon sales export')
    raise InputError('CSV columns not recognized. Use date,instrument,close or an Amazon export with OrderDate,ProductID,UnitPrice.')

def _parse_day(value):
    if not value: return None
    s=str(value).strip()
    try: return date.fromisoformat(s[:10]).isoformat()
    except ValueError: return None

def _number(value):
    try:
        n=float(value)
        return n if n == n and abs(n) != float('inf') else None
    except (TypeError,ValueError): return None

def _grouped_csv(grouped, source_name):
    out=io.StringIO(); w=csv.writer(out,lineterminator='\n'); w.writerow(['date','instrument','close'])
    for (day,instrument),vals in sorted(grouped.items()):
        if vals: w.writerow([day,instrument,f'{sum(vals)/len(vals):.6f}'.rstrip('0').rstrip('.')])
    content=out.getvalue()
    try: normalized=list(csv.DictReader(io.StringIO(content)))
    except Exception: normalized=[]
    if not normalized: raise InputError(f'No usable date/product/price rows found in {source_name}.')
    return content, {'format':source_name+' → daily average unit price','rows':len(normalized),'instruments':len({r['instrument'] for r in normalized})}

def convert_uploaded_file(filename: str, raw: bytes) -> tuple[str, dict]:
    suffix=Path(filename).suffix.lower()
    if suffix == '.csv':
        try: text=raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            try: text=raw.decode('cp1252')
            except UnicodeDecodeError as exc: raise InputError('CSV encoding must be UTF-8 or Windows-1252.') from exc
        return parse_csv_prices(text)
    if suffix in ('.xlsx', '.xlsm'):
        try:
            from openpyxl import load_workbook
        except ImportError as exc: raise InputError('Excel upload requires openpyxl. Run: pip install -r requirements.txt') from exc
        try:
            wb=load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            ws=wb.active
            iterator=ws.iter_rows(values_only=True)
            headers=[str(v).strip() if v is not None else '' for v in next(iterator)]
            indexes={h:i for i,h in enumerate(headers)}
            required={'StockCode','InvoiceDate','UnitPrice'}
            if not required.issubset(indexes):
                raise InputError('Excel columns not recognized. Expected Online Retail columns: StockCode, InvoiceDate, UnitPrice.')
            grouped=defaultdict(list)
            for row in iterator:
                dayval=row[indexes['InvoiceDate']] if indexes['InvoiceDate'] < len(row) else None
                skuval=row[indexes['StockCode']] if indexes['StockCode'] < len(row) else None
                priceval=row[indexes['UnitPrice']] if indexes['UnitPrice'] < len(row) else None
                if isinstance(dayval, datetime): day=dayval.date().isoformat()
                elif isinstance(dayval, date): day=dayval.isoformat()
                else: day=_parse_day(dayval)
                sku=str(skuval).strip() if skuval is not None else ''
                price=_number(priceval)
                if day and sku and price is not None and price > 0: grouped[(day,'OR-'+sku)].append(price)
            wb.close()
            return _grouped_csv(grouped, 'Online Retail Excel')
        except InputError: raise
        except Exception as exc: raise InputError(f'Could not read Excel workbook: {exc}') from exc
    raise InputError('Supported upload formats: .csv and .xlsx')

def suggest_rules(csv_text: str) -> list[dict]:
    """Create editable starter rules for the instruments in the active dataset."""
    grouped=defaultdict(list)
    try:
        for row in csv.DictReader(io.StringIO(csv_text)):
            day=_parse_day(row.get('date')); instrument=(row.get('instrument') or '').strip(); close=_number(row.get('close'))
            if day and instrument and close is not None: grouped[instrument].append((day,close))
    except Exception:
        return []
    ranked=sorted(grouped.items(),key=lambda item:(-len(item[1]),item[0]))
    rules=[]
    for idx,(instrument,series) in enumerate(ranked[:2],start=1):
        values=sorted(v for _,v in series)
        if not values: continue
        mid=len(values)//2
        threshold=values[mid] if len(values)%2 else (values[mid-1]+values[mid])/2
        rules.append({'id':f'auto_level_{idx}','instrument':instrument,'type':'level','direction':'below','level':round(threshold,6)})
    if ranked and len(ranked[0][1])>=20:
        rules.append({'id':'auto_ma_1','instrument':ranked[0][0],'type':'ma_cross','short':5,'long':20,'direction':'above'})
    return rules

def auth_user(handler):
    header=handler.headers.get('Authorization','')
    if not header.startswith('Bearer '): return None
    return verify_token(header[7:].strip())

class Handler(BaseHTTPRequestHandler):
    def _send(self, status, body, content_type='application/json; charset=utf-8'):
        if isinstance(body,str): body=body.encode('utf-8')
        self.send_response(status); self.send_header('Content-Type',content_type); self.send_header('Content-Length',str(len(body)))
        self.send_header('Cache-Control','no-store'); self.send_header('X-Content-Type-Options','nosniff'); self.end_headers(); self.wfile.write(body)
    def _json(self,status,obj): return self._send(status,json.dumps(obj,ensure_ascii=False), 'application/json; charset=utf-8')
    def _body(self, limit=50_000_000):
        length=int(self.headers.get('Content-Length','0'))
        if length<=0: raise InputError('Request body is empty.')
        if length>limit: raise InputError(f'Upload is too large. Maximum size is {limit//1_000_000} MB.')
        return self.rfile.read(length)
    def _require_auth(self):
        if not auth_user(self): self._json(401,{'error':'Your admin session expired. Sign in again.'}); return False
        return True
    def do_GET(self):
        path=urlparse(self.path).path
        if path in ('/','/index.html'): return self._send(200,FRONTEND.read_text(encoding='utf-8'),'text/html; charset=utf-8')
        if path=='/api/health': return self._json(200,{'status':'ok','engine':'ready','auth':'JWT enabled'})
        if path=='/api/sample':
            if not self._require_auth(): return
            return self._json(200,{'prices_csv':SAMPLE_PRICES.read_text(encoding='utf-8'),'rules':json.loads(SAMPLE_RULES.read_text(encoding='utf-8'))})
        if path=='/api/dataset':
            if not self._require_auth(): return
            name=parse_qs(urlparse(self.path).query).get('name',[''])[0]
            file=DATASETS.get(name)
            if not file or not file.exists(): return self._json(404,{'error':'Dataset not found.'})
            dataset_id=secrets.token_urlsafe(18); content=file.read_text(encoding='utf-8')
            ACTIVE_DATASETS[dataset_id]={'csv':content,'name':{'amazon':'Amazon(1).csv','online-retail':'Online Retail(4).xlsx'}[name],'meta':{'format':'Packaged processed dataset','rows':sum(1 for _ in content.splitlines())-1,'instruments':len({r.split(',')[1] for r in content.splitlines()[1:] if ',' in r})}}
            return self._json(200,{'dataset_id':dataset_id,'name':ACTIVE_DATASETS[dataset_id]['name'],'meta':ACTIVE_DATASETS[dataset_id]['meta'],'suggested_rules':suggest_rules(content)})
        return self._json(404,{'error':'Not found'})
    def do_POST(self):
        path=urlparse(self.path).path
        try:
            if path=='/api/login':
                payload=json.loads(self._body(100_000).decode('utf-8'))
                username=str(payload.get('username','')); password=str(payload.get('password',''))
                # Constant-time comparisons avoid a direct timing leak for the credentials.
                ok=hmac.compare_digest(username,ADMIN_USERNAME) and hmac.compare_digest(password,ADMIN_PASSWORD)
                if not ok: return self._json(401,{'error':'Incorrect admin username or password.'})
                return self._json(200,{'token':make_token(username),'username':username,'role':'admin','expires_in':TOKEN_TTL})
            if path not in ('/api/upload','/api/analyze','/api/copilot'):
                return self._json(404,{'error':'Not found'})
            if not self._require_auth(): return
            if path=='/api/upload':
                raw=self._body(50_000_000)
                content_type=self.headers.get('Content-Type','')
                if not content_type.lower().startswith('multipart/form-data;'):
                    return self._json(400,{'error':'Upload must use multipart/form-data.'})
                msg=BytesParser(policy=policy.default).parsebytes(b'Content-Type: '+content_type.encode()+b'\r\nMIME-Version: 1.0\r\n\r\n'+raw)
                part=next((p for p in msg.iter_parts() if p.get_param('name',header='content-disposition')=='file'),None)
                if part is None: return self._json(400,{'error':'Choose a CSV or XLSX file first.'})
                filename=Path(part.get_filename() or 'upload.csv').name
                filebytes=part.get_payload(decode=True) or b''
                normalized,meta=convert_uploaded_file(filename,filebytes)
                dataset_id=secrets.token_urlsafe(18)
                ACTIVE_DATASETS[dataset_id]={'csv':normalized,'name':filename,'meta':meta}
                return self._json(200,{'dataset_id':dataset_id,'name':filename,'meta':meta,'suggested_rules':suggest_rules(normalized),'message':'Dataset uploaded and mapped into daily instrument prices.'})
            payload=json.loads(self._body(50_000_000).decode('utf-8'))
            dataset_id=payload.get('dataset_id')
            if dataset_id:
                dataset=ACTIVE_DATASETS.get(dataset_id)
                if not dataset: return self._json(400,{'error':'Uploaded dataset expired after a server restart. Upload it again.'})
                prices_csv=dataset['csv']
            else: prices_csv=payload.get('prices_csv','')
            rules_payload=payload.get('rules',[])
            if path=='/api/analyze':
                if not isinstance(prices_csv,str) or not prices_csv.strip(): raise InputError('Upload a price file or load a dataset first.')
                if isinstance(rules_payload,str):
                    try: rules_payload=json.loads(rules_payload)
                    except json.JSONDecodeError as exc: raise InputError(f'Rules are not valid JSON: {exc.msg}.') from exc
                import tempfile
                with tempfile.TemporaryDirectory() as tmp:
                    p=Path(tmp)/'prices.csv'; r=Path(tmp)/'rules.json'; p.write_text(prices_csv,encoding='utf-8'); r.write_text(json.dumps(rules_payload),encoding='utf-8')
                    prices=read_prices(p); rules=read_rules(r)
                alerts=find_alerts(prices,rules)
                instruments=sorted({x.instrument for x in prices})
                result={'alerts':[{'date':a.day.isoformat(),'rule_id':a.rule_id,'instrument':a.instrument,'description':a.description,'line':format_alert(a)} for a in alerts],
                        'summary':{'rows':len(prices),'instruments':len(instruments),'rules':len(rules),'alerts':len(alerts)},'instruments':instruments,'status':'success'}
                return self._json(200,result)
            question=str(payload.get('question','')).strip()[:2000]
            if not question: raise InputError('Ask the copilot a question first.')
            result=payload.get('result') or {}
            prompt_context={'question':question,'summary':result.get('summary',{}),'alerts':(result.get('alerts') or [])[:30],'dataset':ACTIVE_DATASETS.get(dataset_id,{}).get('name','current workspace')}
            api_key=os.getenv('GROQ_API_KEY','').strip()
            if api_key:
                req_body=json.dumps({'model':os.getenv('GROQ_MODEL','llama-3.3-70b-versatile'),'messages':[{'role':'system','content':'You are PricePilot AI Copilot for a price-alert assessment. Explain results clearly and conservatively. Do not claim forecasts or live market data. Use only the supplied context; distinguish facts from suggestions. Keep answers concise and practical.'},{'role':'user','content':json.dumps(prompt_context,ensure_ascii=False)}],'temperature':0.2,'max_tokens':500}).encode()
                req=urllib.request.Request('https://api.groq.com/openai/v1/chat/completions',data=req_body,headers={'Authorization':'Bearer '+api_key,'Content-Type':'application/json'},method='POST')
                try:
                    with urllib.request.urlopen(req,timeout=20) as response: answer=json.loads(response.read().decode())['choices'][0]['message']['content']
                    return self._json(200,{'answer':answer,'mode':'Groq AI','configured':True})
                except Exception as exc:
                    return self._json(200,{'answer':self._local_copilot(question,result)+'\n\nNote: Groq could not be reached; showing local fallback analysis.','mode':'Local fallback','configured':True,'detail':str(exc)[:160]})
            return self._json(200,{'answer':self._local_copilot(question,result)+'\n\nFor generative answers, set GROQ_API_KEY in your environment and restart the app. The local copilot works without an API key.','mode':'Local copilot','configured':False})
        except (InputError,ValueError,TypeError,json.JSONDecodeError) as exc: return self._json(400,{'error':str(exc)})
        except Exception as exc: return self._json(500,{'error':f'Server error: {exc}'})
    def _local_copilot(self, question, result):
        s=result.get('summary') or {}; alerts=result.get('alerts') or []
        if not s: return 'Load a dataset and run alert analysis first. I can then explain alert counts, triggered rules, data quality, and next verification steps.'
        if alerts:
            first=alerts[0]
            details='; '.join(f"{a.get('date')} {a.get('rule_id')} {a.get('instrument')}: {a.get('description')}" for a in alerts[:5])
            return (f"Analysis summary: {s.get('rows',0)} daily price rows across {s.get('instruments',0)} instruments, checked against {s.get('rules',0)} rules. "
                    f"{s.get('alerts',0)} alert(s) fired. Examples: {details}. "
                    "Next: inspect each triggering date and its previous close (for level rules) or both moving averages (for MA rules). This is a rule-based signal, not a buy/sell recommendation.")
        return (f"Analysis summary: {s.get('rows',0)} daily price rows across {s.get('instruments',0)} instruments; {s.get('rules',0)} rules checked and no alerts fired. "
                "That can be a valid result: no configured crossing occurred, the rules may reference instruments absent from the dataset, or moving-average windows may need more observations. Check instrument IDs and rule windows before changing thresholds.")
    def log_message(self,fmt,*args): print('[PricePilot]',fmt%args)

if __name__=='__main__':
    host=os.getenv('HOST','127.0.0.1'); port=int(os.getenv('PORT','8000'))
    print(f'PricePilot AI — Wissen workspace: http://{host}:{port}')
    print(f'Demo admin login: {ADMIN_USERNAME} / {ADMIN_PASSWORD}')
    if JWT_SECRET.startswith('CHANGE-ME'): print('WARNING: Set JWT_SECRET before exposing this server beyond localhost.')
    try: ThreadingHTTPServer((host,port),Handler).serve_forever()
    except KeyboardInterrupt: print('\nServer stopped.')
